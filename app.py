"""
Vendor Contract Generator - Streamlit Web Application
Automates the generation of customized vendor contracts from Excel data and a Word (.docx) template.
Includes support for Word (.docx), PDF (.pdf), both formats, and signature image embedding.
"""

import io
import os
import re
import tempfile
import time
import zipfile
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple

import pandas as pd
import streamlit as st
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches

from pathlib import Path
import subprocess
import shutil
 
 
def convert_to_pdf(docx_path, pdf_path):
    if shutil.which("libreoffice") is None:
        raise RuntimeError(
            "LibreOffice is not installed. "
            "Please check packages.txt and restart the app."
        )
 
    output_dir = os.path.dirname(pdf_path) or "."
    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
 
    try:
        result = subprocess.run(
            [
                "libreoffice",
                f"-env:UserInstallation={Path(profile_dir).resolve().as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output_dir,
                docx_path,
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
 
        generated_pdf = os.path.join(
            output_dir,
            os.path.splitext(os.path.basename(docx_path))[0] + ".pdf",
        )
 
        if result.returncode != 0 or not os.path.exists(generated_pdf):
            raise RuntimeError(
                f"PDF conversion failed: {result.stderr or result.stdout}"
            )
 
        if os.path.abspath(generated_pdf) != os.path.abspath(pdf_path):
            os.replace(generated_pdf, pdf_path)
 
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
 
 
DOCX2PDF_AVAILABLE = shutil.which("libreoffice") is not None


# -----------------------------------------------------------------------------
# Page Configuration & Theming
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Vendor Contract Generator",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Custom CSS for modern enterprise feel
st.markdown(
    """
    <style>
    .main .block-container {
        padding-top: 2rem;
        padding-bottom: 3rem;
        max-width: 1200px;
    }
    .stat-badge {
        display: inline-flex;
        align-items: center;
        padding: 0.35rem 0.85rem;
        border-radius: 9999px;
        font-size: 0.85rem;
        font-weight: 600;
        background-color: #f0fdf4;
        color: #166534;
        border: 1px solid #bbf7d0;
    }
    .tag-chip {
        display: inline-block;
        background: #f1f5f9;
        color: #334155;
        padding: 0.2rem 0.5rem;
        border-radius: 4px;
        font-family: monospace;
        font-size: 0.82rem;
        margin: 0.15rem;
        border: 1px solid #e2e8f0;
    }
    .guide-box {
        background-color: #f8fafc;
        border: 1px solid #e2e8f0;
        border-radius: 8px;
        padding: 1.25rem;
        margin-bottom: 1.5rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# -----------------------------------------------------------------------------
# Helper Functions
# -----------------------------------------------------------------------------
def sanitize_value(val: Any) -> Any:
    """
    Sanitizes values extracted from Pandas DataFrames.
    Converts NaN, None, NaT, and float('nan') into clean empty strings,
    and formats dates nicely.
    """
    if pd.isna(val) or val is None:
        return ""
    if isinstance(val, (datetime, pd.Timestamp)):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, float) and val.is_integer():
        return int(val)
    return val


def sanitize_row_context(row: pd.Series) -> Dict[str, Any]:
    """
    Converts a DataFrame row into a clean dictionary context for Jinja2 / docxtpl.
    Both original column names and slugified versions (lower, snake_case) are provided.
    """
    context: Dict[str, Any] = {}
    for col_name, value in row.items():
        clean_val = sanitize_value(value)
        context[str(col_name)] = clean_val
        
        normalized_key = re.sub(r"[^\w]+", "_", str(col_name).strip().lower()).strip("_")
        if normalized_key and normalized_key not in context:
            context[normalized_key] = clean_val
            
    return context


def auto_detect_columns(columns: List[str]) -> Tuple[Optional[str], Optional[str]]:
    """Heuristically auto-detects vendor ID and vendor Name columns."""
    id_candidates = ["vendor_id", "vendor id", "id", "supplier_id", "vendor_code", "code", "account_id"]
    name_candidates = ["vendor_name", "vendor name", "company_name", "company", "name", "supplier_name", "vendor", "supplier"]

    detected_id = None
    detected_name = None

    for col in columns:
        cleaned = col.strip().lower()
        if cleaned in id_candidates or any(cand in cleaned for cand in ["vendor_id", "vendor id", "supplier_id"]):
            detected_id = col
            break
    if not detected_id:
        for col in columns:
            if "id" in col.strip().lower() or "code" in col.strip().lower():
                detected_id = col
                break

    for col in columns:
        cleaned = col.strip().lower()
        if cleaned in name_candidates or any(cand in cleaned for cand in ["vendor_name", "company"]):
            detected_name = col
            break
    if not detected_name:
        for col in columns:
            if "name" in col.strip().lower() and col != detected_id:
                detected_name = col
                break

    if not detected_id and len(columns) > 0:
        detected_id = columns[0]
    if not detected_name and len(columns) > 1:
        detected_name = columns[1] if columns[1] != detected_id else columns[0]
    elif not detected_name and len(columns) > 0:
        detected_name = columns[0]

    return detected_id, detected_name


def generate_sample_excel() -> bytes:
    """Generates an in-memory sample Excel file for quick testing."""
    sample_data = {
        "vendor_id": ["VEND-101", "VEND-102", "VEND-103", "VEND-104", "VEND-105"],
        "vendor_name": [
            "Apex Logistics Solutions LLC",
            "BluePeak Cloud Technologies Inc.",
            "Cascade Industrial Supplies",
            "Delta Precision Engineering",
            "Evergreen Facility Services",
        ],
        "contact_person": ["Sarah Jenkins", "Michael Chang", "Elena Rostova", "David Miller", "Amina Patel"],
        "email": [
            "sarah.j@apexlogistics.com",
            "mchang@bluepeakcloud.io",
            "e.rostova@cascade-ind.com",
            "dmiller@deltaprecision.com",
            "apatel@evergreenfs.com",
        ],
        "phone": ["(555) 234-8901", "(555) 876-5432", "(555) 345-6789", "(555) 901-2345", "(555) 678-1234"],
        "effective_date": ["2026-10-01", "2026-10-15", "2026-11-01", "2026-11-15", "2026-12-01"],
        "contract_term": ["12 Months", "24 Months", "36 Months", "12 Months", "24 Months"],
        "contract_value": [45000, 120000, 32500, 88000, 19500],
        "currency": ["USD", "USD", "USD", "USD", "USD"],
        "payment_terms": ["Net 30", "Net 15", "Net 45", "Net 30", "Net 30"],
        "state_of_incorporation": ["Delaware", "California", "Washington", "Texas", "Illinois"],
        "notes": [
            "Includes priority 24/7 SLA dispatch.",
            "Cloud infrastructure SLA 99.95% tier.",
            None,
            "Certified ISO-9001 precision supplier.",
            None,
        ],
    }
    df = pd.DataFrame(sample_data)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Vendors")
    return buf.getvalue()


def generate_sample_docx_template() -> bytes:
    """Creates a sample Word document (.docx) containing Jinja2 placeholders and signature tag."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("MASTER VENDOR SERVICES AGREEMENT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub.add_run("Contract Reference: {{ vendor_id }}")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.25
    r = p1.add_run('This Master Vendor Agreement ("Agreement") is entered into as of ')
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    r_date = p1.add_run("{{ effective_date }}")
    r_date.font.name = "Arial"
    r_date.font.size = Pt(10.5)
    r_date.font.bold = True

    r2 = p1.add_run(', by and between Enterprise Corporation ("Client") and ')
    r2.font.name = "Arial"
    r2.font.size = Pt(10.5)

    r_vname = p1.add_run("{{ vendor_name }}")
    r_vname.font.name = "Arial"
    r_vname.font.size = Pt(10.5)
    r_vname.font.bold = True

    r3 = p1.add_run(', a corporation organized under the laws of {{ state_of_incorporation }} ("Vendor").')
    r3.font.name = "Arial"
    r3.font.size = Pt(10.5)

    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Vendor Information & Engagement Summary")
    h1_run.font.name = "Arial"
    h1_run.font.size = Pt(13)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(15, 23, 42)

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"

    rows_data = [
        ("Vendor Identification No.", "{{ vendor_id }}"),
        ("Primary Contact Person", "{{ contact_person }}"),
        ("Contact Email / Phone", "{{ email }} / {{ phone }}"),
        ("Engagement Term", "{{ contract_term }}"),
        ("Contract Value & Currency", "{{ contract_value }} {{ currency }}"),
        ("Invoicing / Payment Terms", "{{ payment_terms }}"),
    ]

    for i, (label, val) in enumerate(rows_data):
        row = table.rows[i]
        c0 = row.cells[0].paragraphs[0].add_run(label)
        c0.font.name = "Arial"
        c0.font.size = Pt(9.5)
        c0.font.bold = True
        
        c1 = row.cells[1].paragraphs[0].add_run(val)
        c1.font.name = "Arial"
        c1.font.size = Pt(9.5)

    doc.add_paragraph()
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. Scope Notes & Special Conditions")
    h2_run.font.name = "Arial"
    h2_run.font.size = Pt(13)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(15, 23, 42)

    p_notes = doc.add_paragraph()
    p_notes.paragraph_format.line_spacing = 1.25
    r_notes = p_notes.add_run(
        "{% if notes %}{{ notes }}{% else %}No special scope notes or riders specified for this vendor.{% endif %}"
    )
    r_notes.font.name = "Arial"
    r_notes.font.size = Pt(10.5)

    doc.add_paragraph()
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. Signatures")
    h3_run.font.name = "Arial"
    h3_run.font.size = Pt(13)
    h3_run.font.bold = True

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.style = "Table Grid"

    sig_table.rows[0].cells[0].paragraphs[0].add_run("CLIENT:\nEnterprise Corporation").font.bold = True
    sig_table.rows[0].cells[1].paragraphs[0].add_run("VENDOR:\n{{ vendor_name }}").font.bold = True

    sig_table.rows[1].cells[0].paragraphs[0].add_run(
        "\nAuthorized Signature:\n{{ signature_img }}\nName: Authorized Officer\nDate: {{ effective_date }}"
    )
    sig_table.rows[1].cells[1].paragraphs[0].add_run(
        "\nVendor Signature:\n__________________________\nName: {{ contact_person }}\nDate: __________________________"
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -----------------------------------------------------------------------------
# Main Application Flow
# -----------------------------------------------------------------------------
def main():
    st.title("📄 Vendor Contract Generator")
    st.markdown(
        "Automate the batch generation of personalized vendor contracts with custom corporate signatures."
    )

    if "vendor_multiselect" not in st.session_state:
        st.session_state.vendor_multiselect = []

    # -------------------------------------------------------------------------
    # Sidebar: File Uploads & Configuration
    # -------------------------------------------------------------------------
    with st.sidebar:
        st.header("📂 1. Source Files")

        excel_file = st.file_uploader(
            "Excel Vendor Data (.xlsx, .xls)",
            type=["xlsx", "xls"],
            help="Upload the spreadsheet containing vendor rows and attributes.",
        )

        template_file = st.file_uploader(
            "Word Contract Template (.docx)",
            type=["docx"],
            help="Upload the Word template with Jinja2 placeholders (e.g. {{ vendor_name }}).",
        )

        st.divider()
        st.header("✍️ 2. Signature & Output Settings")

        signature_file = st.file_uploader(
            "Authorized Signature Image (PNG, JPG)",
            type=["png", "jpg", "jpeg"],
            help="Upload an image of your signature to embed into the contract.",
        )

        sig_width_inches = st.slider(
            "Signature Image Width (Inches)",
            min_value=1.0,
            max_value=3.5,
            value=1.75,
            step=0.25,
            help="Adjusts the display size of the signature in the document.",
        )
        
        output_format = st.selectbox(
            "Output Format", 
            options=["Word (.docx)", "PDF (.pdf)", "Both (.docx & .pdf)"],
            help="Choose the format of the generated contracts.",
        )

        st.divider()

        st.markdown("**Need test files to try it out?**")
        col_s1, col_s2 = st.columns(2)
        with col_s1:
            sample_excel_bytes = generate_sample_excel()
            st.download_button(
                "📥 Sample Excel",
                data=sample_excel_bytes,
                file_name="sample_vendor_data.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
            )
        with col_s2:
            try:
                sample_doc_bytes = generate_sample_docx_template()
                st.download_button(
                    "📥 Sample Docx",
                    data=sample_doc_bytes,
                    file_name="sample_contract_template.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    width="stretch",
                )
            except Exception:
                pass

        st.divider()

        with st.expander("ℹ️ Template Placeholder Guide", expanded=False):
            st.markdown(
                """
                In your Word (`.docx`) file, insert variables enclosed in double curly braces:
                
                - `{{ vendor_name }}`
                - `{{ vendor_id }}`
                - `{{ effective_date }}`
                - `{{ contract_value }}`
                - `{{ signature_img }}` *(Inserts your uploaded signature image)*
                """
            )

    # -------------------------------------------------------------------------
    # Step 1: Validate Uploads & Parse Data
    # -------------------------------------------------------------------------
    if not excel_file or not template_file:
        st.info("👈 Please upload both an **Excel file** and a **Word template** in the sidebar to begin.")
        st.markdown(
            """
            ### How it works:
            1. **Upload Data & Template**: Provide your Excel records and Word contract template.
            2. **Upload Signature**: Attach your signature image file in the sidebar.
            3. **Select Vendors**: Choose specific vendors or **Select All** for batch processing.
            4. **Generate & Download**: Download your customized contracts containing your embedded signature!
            """
        )
        return

    try:
        df = pd.read_excel(excel_file)
    except Exception as e:
        st.error(f"❌ Failed to parse Excel file: {str(e)}")
        return

    if df.empty:
        st.warning("⚠️ The uploaded Excel file does not contain any rows.")
        return

    try:
        template_bytes = template_file.getvalue()
        _ = DocxTemplate(io.BytesIO(template_bytes))
    except Exception as e:
        st.error(f"❌ The uploaded Word template is invalid or corrupted: {str(e)}")
        return

    columns_list = [str(c) for c in df.columns.tolist()]
    detected_id_col, detected_name_col = auto_detect_columns(columns_list)

    # -------------------------------------------------------------------------
    # Step 2: Column Identification & Data Preview
    # -------------------------------------------------------------------------
    st.subheader("📊 2. Data Preview & Column Mapping")

    with st.expander("🔍 Inspect Uploaded Vendor Data & Key Column Selection", expanded=True):
        col_c1, col_c2, col_c3 = st.columns([1, 1, 1])

        with col_c1:
            id_col = st.selectbox(
                "Vendor ID Column:",
                options=columns_list,
                index=columns_list.index(detected_id_col) if detected_id_col in columns_list else 0,
            )
        with col_c2:
            name_col = st.selectbox(
                "Vendor Name Column:",
                options=columns_list,
                index=columns_list.index(detected_name_col) if detected_name_col in columns_list else 0,
            )
        with col_c3:
            st.metric("Total Vendors Found", len(df))

        st.dataframe(df.fillna(""), width="stretch", height=220)

        with st.expander("📋 View Available Template Placeholders", expanded=False):
            st.markdown("Use any of these exact tags inside your Word `.docx` template (plus `{{ signature_img }}`):")
            tags_html = "".join([f'<span class="tag-chip">{{{{ {col} }}}}</span>' for col in columns_list])
            tags_html += '<span class="tag-chip">{{ signature_img }}</span>'
            st.markdown(tags_html, unsafe_allow_html=True)

    # -------------------------------------------------------------------------
    # Step 3: Interactive Selection (Select All / Deselect All)
    # -------------------------------------------------------------------------
    st.subheader("🎯 3. Select Vendors for Contract Generation")

    vendor_options: Dict[int, str] = {}
    for idx, row in df.iterrows():
        v_id = str(sanitize_value(row.get(id_col, idx)))
        v_name = str(sanitize_value(row.get(name_col, f"Vendor {idx}")))
        vendor_options[int(idx)] = f"[{v_id}] - {v_name}"

    all_indices = list(vendor_options.keys())

    st.session_state.vendor_multiselect = [
        i for i in st.session_state.vendor_multiselect if i in all_indices
    ]

    ctrl_col1, ctrl_col2, ctrl_col3 = st.columns([1, 1, 3])

    with ctrl_col1:
        if st.button("✅ Select All", width="stretch"):
            st.session_state.vendor_multiselect = all_indices
            st.rerun()

    with ctrl_col2:
        if st.button("❌ Deselect All", width="stretch"):
            st.session_state.vendor_multiselect = []
            st.rerun()

    with ctrl_col3:
        current_count = len(st.session_state.vendor_multiselect)
        st.markdown(
            f"<div style='padding-top: 0.4rem;'>Selected: <b>{current_count}</b> of <b>{len(df)}</b> vendors</div>",
            unsafe_allow_html=True,
        )

    selected_indices = st.multiselect(
        "Choose vendors to generate contracts for:",
        options=all_indices,
        format_func=lambda x: vendor_options.get(x, str(x)),
        key="vendor_multiselect",
    )

    if not selected_indices:
        st.warning("⚠️ Please select at least one vendor to generate contracts.")
        return

    # -------------------------------------------------------------------------
    # Step 4: Generation Logic (docxtpl & signature embedding)
    # -------------------------------------------------------------------------
    st.subheader("⚡ 4. Document Generation")

    generate_clicked = st.button(
        f"Generate {len(selected_indices)} Contract{'s' if len(selected_indices) > 1 else ''}",
        type="primary",
        width="stretch",
    )

    if generate_clicked:
        if output_format in ["PDF (.pdf)", "Both (.docx & .pdf)"] and not DOCX2PDF_AVAILABLE:
            st.error("❌ PDF generation requires LibreOffice installed on the server.")
            return

        progress_bar = st.progress(0)
        status_text = st.empty()

        generated_files: List[Tuple[str, bytes]] = []
        errors: List[str] = []

        total_selected = len(selected_indices)
        start_time = time.time()

        for i, idx in enumerate(selected_indices):
            row = df.loc[idx]
            v_id = str(sanitize_value(row.get(id_col, idx)))
            v_name = str(sanitize_value(row.get(name_col, f"Vendor_{idx}")))
            
            safe_name = re.sub(r"[^\w\-_\. ]", "_", f"{v_id}_{v_name}").strip()
            status_text.text(f"Rendering contract ({i+1}/{total_selected}): {v_name}...")

            try:
                context = sanitize_row_context(row)
                doc = DocxTemplate(io.BytesIO(template_bytes))
                
                # Handle signature image insertion if uploaded
                if signature_file is not None:
                    sig_bytes = signature_file.getvalue()
                    sig_stream = io.BytesIO(sig_bytes)
                    context["signature_img"] = InlineImage(doc, sig_stream, width=Inches(sig_width_inches))
                else:
                    context["signature_img"] = "[Signature Placeholder - No Image Provided]"

                doc.render(context)
                
                # Generate Word (.docx)
                if output_format in ["Word (.docx)", "Both (.docx & .pdf)"]:
                    docx_filename = f"Contract_{safe_name}.docx"
                    out_stream = io.BytesIO()
                    doc.save(out_stream)
                    generated_files.append((docx_filename, out_stream.getvalue()))

                # Generate PDF (.pdf)
                if output_format in ["PDF (.pdf)", "Both (.docx & .pdf)"]:
                    pdf_filename = f"Contract_{safe_name}.pdf"
                    with tempfile.TemporaryDirectory() as tmpdir:
                        temp_docx = os.path.join(tmpdir, "temp.docx")
                        temp_pdf = os.path.join(tmpdir, "temp.pdf")
                        
                        doc.save(temp_docx)
                        convert_to_pdf(temp_docx, temp_pdf)
                        
                        with open(temp_pdf, "rb") as f:
                            pdf_bytes = f.read()
                            
                    generated_files.append((pdf_filename, pdf_bytes))

            except Exception as render_err:
                errors.append(f"Failed to generate for row {idx} ({v_name}): {str(render_err)}")

            progress_bar.progress((i + 1) / total_selected)

        elapsed = time.time() - start_time
        status_text.empty()
        progress_bar.empty()

        if errors:
            for err in errors:
                st.error(f"⚠️ {err}")

        if not generated_files:
            st.error("❌ No contracts could be generated. Check the template placeholders and try again.")
            return

        st.success(
            f"🎉 Successfully rendered **{len(generated_files)}** file{'s' if len(generated_files) > 1 else ''} in **{elapsed:.2f}s**!"
        )

        # ---------------------------------------------------------------------
        # Step 5: Output & Downloads
        # ---------------------------------------------------------------------
        st.subheader("📥 5. Download Contracts")

        if len(generated_files) == 1:
            fname, fbytes = generated_files[0]
            mime_type = "application/pdf" if fname.endswith(".pdf") else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            st.download_button(
                label=f"📄 Download {fname}",
                data=fbytes,
                file_name=fname,
                mime=mime_type,
                type="primary",
            )
        else:
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zip_file:
                for fname, fbytes in generated_files:
                    zip_file.writestr(fname, fbytes)

            zip_bytes = zip_buffer.getvalue()
            timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
            zip_filename = f"Vendor_Contracts_{timestamp_str}.zip"

            col_d1, col_d2 = st.columns([2, 1])
            with col_d1:
                st.download_button(
                    label=f"📦 Download All {len(generated_files)} Contracts as ZIP ({len(zip_bytes) / 1024:.1f} KB)",
                    data=zip_bytes,
                    file_name=zip_filename,
                    mime="application/zip",
                    type="primary",
                    width="stretch",
                )

            with st.expander("📄 Generated Files Manifest", expanded=True):
                summary_data = [
                    {"File Name": fn, "Size (KB)": round(len(fb) / 1024, 2)}
                    for fn, fb in generated_files
                ]
                st.table(pd.DataFrame(summary_data))


if __name__ == "__main__":
    main()